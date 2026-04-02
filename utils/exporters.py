import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from tkinter import filedialog, messagebox
from datetime import datetime


class Exporters:

    @staticmethod
    def exportar_excel(datos, columnas, titulo):
        filename = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"{titulo}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        if not filename:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = titulo

        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

        # Encabezados
        for col, header in enumerate(columnas, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill

        # Datos
        for row_idx, row in enumerate(datos, 2):
            for col_idx, col_name in enumerate(columnas, 1):
                key = col_name.lower().replace(' ', '_')
                valor = row.get(key, '')
                ws.cell(row=row_idx, column=col_idx, value=valor)

        # Ajustar columnas
        for col in ws.columns:
            max_len = 0
            for cell in col:
                try:
                    max_len = max(max_len, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

        wb.save(filename)
        messagebox.showinfo("Éxito", f"Exportado a:\n{filename}")

    @staticmethod
    def exportar_pdf(datos, columnas, titulo):
        filename = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF files", "*.pdf")],
            initialfile=f"{titulo}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        )
        if not filename:
            return

        doc = SimpleDocTemplate(filename, pagesize=landscape(letter))
        elementos = []
        styles = getSampleStyleSheet()

        # Título
        elementos.append(Paragraph(titulo, styles['Title']))
        elementos.append(Spacer(1, 0.2 * inch))
        elementos.append(Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
        elementos.append(Spacer(1, 0.2 * inch))

        # Tabla
        data = [columnas]
        for row in datos:
            fila = []
            for col in columnas:
                key = col.lower().replace(' ', '_')
                valor = row.get(key, '')
                if col == 'Precio' and valor:
                    fila.append(f"${float(valor):,.2f}")
                else:
                    fila.append(str(valor))
            data.append(fila)

        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        elementos.append(table)

        doc.build(elementos)
        messagebox.showinfo("Éxito", f"PDF exportado:\n{filename}")